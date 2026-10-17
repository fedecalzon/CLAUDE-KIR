import openpyxl,sys,json,warnings,datetime
warnings.filterwarnings("ignore")
wbv=openpyxl.load_workbook(sys.argv[1],data_only=True)
wbf=openpyxl.load_workbook(sys.argv[1])
out=[]
for ws in wbv.worksheets:
    wf=wbf[ws.title]
    rows=[]
    for r in range(1,ws.max_row+1):
        row=[]
        for c in range(1,ws.max_column+1):
            v=ws.cell(r,c).value; f=wf.cell(r,c).value
            if isinstance(v,datetime.datetime): v={"$d":v.strftime("%Y-%m-%d")}
            if isinstance(f,str) and f.startswith("="): row.append({"$f":f,"v":v})
            else: row.append(v)
        rows.append(row)
    out.append({"name":ws.title,"hidden":ws.sheet_state!="visible","values":rows,
      "heights":{r:ws.row_dimensions[r].height for r in range(1,ws.max_row+1) if ws.row_dimensions[r].height},
      "merges":[str(m) for m in ws.merged_cells.ranges]})
json.dump(out,open(sys.argv[2],"w"))

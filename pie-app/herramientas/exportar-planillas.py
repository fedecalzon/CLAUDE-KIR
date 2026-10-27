"""Exporta valores cacheados de XLSX para auditar el lector de Apps Script sin acceder a Google.

Requiere openpyxl. El JSON contiene datos de obra: conservar en almacenamiento privado.
No modifica los libros ni calcula fórmulas; utiliza sus últimos valores guardados.
"""
import argparse
import datetime
import json
from pathlib import Path
from openpyxl import load_workbook

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--pie', required=True)
    ap.add_argument('--fod', required=True)
    ap.add_argument('--fog', required=True)
    ap.add_argument('--salida', required=True)
    args = ap.parse_args()
    def value(v):
        if isinstance(v, (datetime.datetime, datetime.date)):
            return v.strftime('%Y-%m-%d')
        return '' if v is None else v
    out = {}
    for key in ('PIE', 'FOD', 'FOG'):
        book = load_workbook(getattr(args, key.lower()), read_only=True, data_only=True, keep_links=False)
        try:
            out[key] = {}
            for sheet in book:
                rows = [[value(v) for v in row] for row in sheet.iter_rows(max_col=60 if key == 'FOG' else 30, values_only=True)]
                while rows and all(v == '' for v in rows[-1]):
                    rows.pop()
                out[key][sheet.title] = rows
        finally:
            book.close()
    target = Path(args.salida)
    # No sobreescribir sin querer otro archivo ni exportaciones previas.
    with target.open('x', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False)
    print('Exportación creada; contiene datos privados de los libros indicados.')

if __name__ == '__main__':
    main()
